import base64
import binascii
import hashlib
import os
import statistics
import sys
import time
from pathlib import Path

import cv2
import numpy as np
import pywt
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
GEO_DIR = ROOT / "All_geo_attacked"
DECODED_DIR = ROOT / "All_geo_decoded"
RESULT_PATH = ROOT / "All_methods_geo_results.xlsx"
PRIVATE_KEY_PATH = ROOT / "private_key.pem"
PUBLIC_KEY_PATH = ROOT / "public_key.pem"

WATERMARK_TEXT = "ASys Encryption"
DCT_Q = 8.0
DWT_Q = 8.0
SS_THRESHOLD = 0.25
METHODS = ("LSB", "DCT", "DWT", "SpatialSS", "DWTSS", "TrustMark")


def load_and_validate_rsa():
    if not PRIVATE_KEY_PATH.exists() or not PUBLIC_KEY_PATH.exists():
        raise FileNotFoundError("private_key.pem and public_key.pem must be beside this script")
    private = serialization.load_pem_private_key(PRIVATE_KEY_PATH.read_bytes(), password=None)
    public = serialization.load_pem_public_key(PUBLIC_KEY_PATH.read_bytes())
    if private.key_size != 2048 or public.key_size != 2048:
        raise ValueError(f"RSA-2048 required; got private={private.key_size}, public={public.key_size}")
    signature = private.sign(WATERMARK_TEXT.encode(), padding.PKCS1v15(), hashes.SHA256())
    public.verify(signature, WATERMARK_TEXT.encode(), padding.PKCS1v15(), hashes.SHA256())
    if len(signature) != 256:
        raise ValueError(f"RSA-2048 signature must be 256 bytes; got {len(signature)}")
    token = base64.b32encode(hashlib.sha256(signature).digest()).decode("ascii")[:8]
    return public, signature, token


def read_bgr(path):
    img = cv2.imread(str(path), cv2.IMREAD_UNCHANGED)
    if img is None:
        return None
    if img.ndim == 2:
        return cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
    if img.shape[2] == 4:
        return cv2.cvtColor(img, cv2.COLOR_BGRA2BGR)
    return img[:, :, :3]


def ref_bits(signature):
    return np.unpackbits(np.frombuffer(signature, dtype=np.uint8)).astype(np.uint8)


def pack_bits(data):
    usable = data[:data.size // 8 * 8]
    return np.packbits(usable).tobytes()


def verify(public, recovered):
    try:
        public.verify(recovered, WATERMARK_TEXT.encode(), padding.PKCS1v15(), hashes.SHA256())
        return True
    except Exception:
        return False


def bit_error_rate(recovered, reference):
    if recovered.size == 0:
        return 1.0
    n = min(recovered.size, reference.size)
    errors = int(np.count_nonzero(recovered[:n] != reference[:n]))
    errors += reference.size - n
    return errors / reference.size


def decode_lsb(img, expected_len):
    flat = img.reshape(-1, 3)
    if flat.shape[0] < 32:
        return np.empty(0, dtype=np.uint8), "image has fewer than 32 pixels"
    header = "".join(str(int(v & 1)) for v in flat[:32, 0])
    declared = int(header, 2)
    capacity = flat.shape[0] - 32
    if declared <= 0 or declared > capacity or declared > 10_000_000:
        return np.empty(0, dtype=np.uint8), f"invalid LSB header length {declared} (capacity {capacity})"
    take = min(declared, expected_len)
    return (flat[32:32 + take, 0] & 1).astype(np.uint8), ""


def decode_dct(img, expected_len):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    h, w = gray.shape
    positions = [(4, 4), (3, 3), (2, 2), (5, 5)]
    capacity = (h // 8) * (w // 8) * len(positions)
    if capacity < expected_len:
        return np.empty(0, dtype=np.uint8), f"DCT capacity {capacity} < {expected_len}"
    out = []
    for y in range(0, h - 7, 8):
        for x in range(0, w - 7, 8):
            block = cv2.dct(gray[y:y + 8, x:x + 8].astype(np.float32))
            for ry, rx in positions:
                out.append(int(np.round(block[ry, rx] / DCT_Q)) & 1)
                if len(out) == expected_len:
                    return np.asarray(out, dtype=np.uint8), ""
    return np.asarray(out, dtype=np.uint8), "short DCT decode"


def decode_dwt(img, expected_len):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY).astype(np.float32)
    _, (lh, _, _) = pywt.dwt2(gray, "haar")
    flat = lh.ravel()
    if flat.size < expected_len:
        return np.empty(0, dtype=np.uint8), f"DWT capacity {flat.size} < {expected_len}"
    out = np.array([int(np.round(v / DWT_Q)) & 1 for v in flat[:expected_len]], dtype=np.uint8)
    return out, ""


def seed_base():
    return int(hashlib.sha256(WATERMARK_TEXT.encode()).hexdigest(), 16) % (2**31 - 1)


def spread_detect(flat, expected_len):
    if flat.size < expected_len:
        return np.empty(0, dtype=np.uint8), f"spread capacity {flat.size} < {expected_len}"
    chunk = flat.size // expected_len
    out = np.zeros(expected_len, dtype=np.uint8)
    base = seed_base()
    for i in range(expected_len):
        start = i * chunk
        end = flat.size if i == expected_len - 1 else (i + 1) * chunk
        pn = np.random.RandomState(base + i).choice([-1, 1], end - start).astype(np.float32)
        out[i] = 1 if float(np.sum(flat[start:end] * pn)) >= 0 else 0
    return out, ""


def decode_spatial_ss(img, expected_len):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY).astype(np.float32)
    return spread_detect(gray.ravel(), expected_len)


def decode_dwt_ss(img, expected_len):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY).astype(np.float32)
    _, (lh, _, _) = pywt.dwt2(gray, "haar")
    return spread_detect(lh.ravel(), expected_len)


def get_trustmark():
    try:
        from trustmark import TrustMark
    except ImportError as exc:
        raise RuntimeError("TrustMark is not installed. Run: pip install trustmark") from exc
    try:
        return TrustMark(verbose=False, model_type="Q", use_ECC=True)
    except TypeError:
        return TrustMark(verbose=False, model_type="Q")


def trustmark_decode(tm, img):
    from PIL import Image
    pil = Image.fromarray(cv2.cvtColor(img, cv2.COLOR_BGR2RGB))
    result = tm.decode(pil)
    if isinstance(result, tuple):
        secret = result[0]
        present = bool(result[1]) if len(result) > 1 else bool(secret)
    else:
        secret, present = result, bool(result)
    if isinstance(secret, bytes):
        secret = secret.decode("utf-8", errors="replace")
    return str(secret or "").strip(), present


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(55, max(12, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width


def main():
    start = time.time()
    public, signature, expected_token = load_and_validate_rsa()
    reference = ref_bits(signature)
    if not GEO_DIR.is_dir():
        print("Attack folder not found:", GEO_DIR)
        return 1
    folders = sorted({p.parent for p in GEO_DIR.rglob("*_geo.png")})
    if not folders:
        print("No *_geo.png images found in:", GEO_DIR)
        return 1
    DECODED_DIR.mkdir(parents=True, exist_ok=True)
    tm = get_trustmark()
    rows = []

    print("=" * 72)
    print("ALL METHODS: GEOMETRIC ATTACK DETECTION")
    print("Folders           :", len(folders))
    print("RSA               : 2048 bits / 256-byte signature")
    print("Expected TrustMark:", expected_token)
    print("=" * 72)

    decoders = {
        "LSB": decode_lsb,
        "DCT": decode_dct,
        "DWT": decode_dwt,
        "SpatialSS": decode_spatial_ss,
        "DWTSS": decode_dwt_ss,
    }
    for index, folder in enumerate(folders, 1):
        item = str(folder.relative_to(GEO_DIR))
        print(f"[{index}/{len(folders)}] {item}")
        text_dir = DECODED_DIR / item
        text_dir.mkdir(parents=True, exist_ok=True)
        for method in METHODS:
            path = folder / f"{method}_geo.png"
            if not path.exists():
                rows.append([item, method, False, False, False, None, "", "missing file"])
                continue
            img = read_bgr(path)
            if img is None:
                rows.append([item, method, False, False, False, None, "", "cannot read image"])
                continue
            try:
                if method == "TrustMark":
                    recovered_text, present = trustmark_decode(tm, img)
                    exact = recovered_text == expected_token
                    rows.append([item, method, present, exact, exact, None, recovered_text, ""])
                    (text_dir / f"{method}.txt").write_text(
                        f"Present: {present}\nExact: {exact}\nExpected: {expected_token}\nRecovered: {recovered_text}\n",
                        encoding="utf-8",
                    )
                    continue
                recovered_bits, error = decoders[method](img, reference.size)
                ber = bit_error_rate(recovered_bits, reference)
                exact = recovered_bits.size == reference.size and bool(np.array_equal(recovered_bits, reference))
                recovered_bytes = pack_bits(recovered_bits)
                rsa_valid = recovered_bits.size == reference.size and verify(public, recovered_bytes)
                detected = (ber < SS_THRESHOLD) if method in ("SpatialSS", "DWTSS") else rsa_valid
                rows.append([item, method, detected, exact, rsa_valid, ber,
                             binascii.hexlify(recovered_bytes).decode("ascii"), error])
                (text_dir / f"{method}.txt").write_text(
                    f"Detected: {detected}\nExact: {exact}\nRSA valid: {rsa_valid}\nBER: {ber:.8f}\nError: {error}\n"
                    f"Recovered hex: {binascii.hexlify(recovered_bytes).decode('ascii')}\n",
                    encoding="utf-8",
                )
            except Exception as exc:
                rows.append([item, method, False, False, False, None, "", str(exc)])
                print(f"  [FAIL] {method}: {exc}")

    wb = Workbook()
    detail = wb.active
    detail.title = "Detailed Results"
    detail.append(["Image", "Method", "Detected", "Exact", "RSA valid", "BER", "Recovered", "Error"])
    for row in rows:
        detail.append(row)
    style_sheet(detail)

    summary = wb.create_sheet("Summary")
    summary.append(["Method", "N", "Detected", "Detection rate", "Exact", "Exact rate",
                    "RSA valid", "RSA rate", "Mean BER", "BER std", "Failed/missing"])
    for method in METHODS:
        selected = [r for r in rows if r[1] == method]
        valid_ber = [float(r[5]) for r in selected if r[5] is not None]
        n = len(selected)
        detected = sum(bool(r[2]) for r in selected)
        exact = sum(bool(r[3]) for r in selected)
        rsa_ok = sum(bool(r[4]) for r in selected)
        failed = sum(bool(r[7]) for r in selected)
        summary.append([
            method, n, detected, detected / n if n else 0,
            exact, exact / n if n else 0,
            rsa_ok, rsa_ok / n if n else 0,
            statistics.fmean(valid_ber) if valid_ber else None,
            statistics.pstdev(valid_ber) if len(valid_ber) > 1 else (0.0 if valid_ber else None),
            failed,
        ])
    style_sheet(summary)
    for row in range(2, summary.max_row + 1):
        for col in (4, 6, 8, 9, 10):
            summary.cell(row, col).number_format = "0.000000"
    wb.save(RESULT_PATH)

    print("=" * 72)
    for method in METHODS:
        selected = [r for r in rows if r[1] == method]
        n = len(selected)
        detected = sum(bool(r[2]) for r in selected)
        exact = sum(bool(r[3]) for r in selected)
        print(f"{method:<12} N={n:<5} Detect={detected / n if n else 0:.4f} Exact={exact / n if n else 0:.4f}")
    print("Runtime           : %.2f seconds" % (time.time() - start))
    print("Excel result      :", RESULT_PATH)
    print("Decoded output    :", DECODED_DIR)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print("[FATAL]", type(exc).__name__ + ":", exc)
        sys.exit(1)
